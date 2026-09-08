#!/usr/bin/env python3
"""
KMTC (ekmtc.com) sailing schedule exporter.

Calls the same public schedule API that KMTC's own website
(https://www.ekmtc.com/index.html#/schedule/leg) uses to fetch the outbound
vessel schedule between two ports, for a range of months, and writes the
combined, de-duplicated result to a formatted Excel (.xlsx) file.

Default route : LAEM CHABANG, Thailand (LCH) -> SHANGHAI, China (SHA)
Default period: current month through December of the current year

Usage:
    python kmtc_schedule_export.py
    python kmtc_schedule_export.py --pol LCH --pol-name "LAEM CHABANG" --pol-ctr TH \
                                    --pod SHA --pod-name SHANGHAI --pod-ctr CN
    python kmtc_schedule_export.py --start 2026-09 --end 2026-12 --out my_schedule.xlsx

Requires: requests, openpyxl
    pip install requests openpyxl
"""

import argparse
import datetime as dt
from typing import Dict, List, Tuple

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

API_URL = "https://api.ekmtc.com/schedule/schedule/leg/search-schedule"

# A normal browser User-Agent + the site's own Origin/Referer keep the request
# looking like it came from the KMTC website itself.
REQUEST_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    ),
    "Referer": "https://www.ekmtc.com/",
    "Origin": "https://www.ekmtc.com",
    "Accept": "application/json, text/plain, */*",
}


def month_range(start: dt.date, end: dt.date):
    """Yield (year, month) tuples from start's month through end's month, inclusive."""
    y, m = start.year, start.month
    while (y, m) <= (end.year, end.month):
        yield y, m
        m += 1
        if m == 13:
            m, y = 1, y + 1


def fetch_month(
    session: requests.Session,
    pol_ctr: str, pol: str, pol_name: str,
    pod_ctr: str, pod: str, pod_name: str,
    year: int, month: int,
) -> List[Dict]:
    """Call the KMTC schedule API for one calendar month; return the raw sailing list."""
    params = {
        "startPlcCd": pol,
        "searchMonth": f"{month:02d}",
        "bound": "O",
        "startPlcName": pol_name,
        "destPlcCd": pod,
        "searchYear": str(year),
        "startCtrCd": pol_ctr,
        "destCtrCd": pod_ctr,
        "destPlcName": pod_name,
        "main": "N",
        "legIdx": "0",
        "vslType01": "01",
        "vslType03": "03",
        "eiCatCd": "O",          # O = export/outbound
        "calendarOrList": "C",   # C = calendar view (matches the website's default)
        "filterTs": "Y",
        "filterDirect": "Y",
    }
    resp = session.get(API_URL, params=params, headers=REQUEST_HEADERS, timeout=20)
    resp.raise_for_status()
    data = resp.json()
    return data.get("listSchedule", [])


def collect_schedule(
    pol_ctr: str, pol: str, pol_name: str,
    pod_ctr: str, pod: str, pod_name: str,
    start: dt.date, end: dt.date,
) -> List[Dict]:
    """Fetch every month in [start, end], de-duplicate sailings that show up in more
    than one month's response (KMTC's calendar view overlaps slightly at month
    boundaries), and keep only sailings whose ETD actually falls in the requested
    date window."""
    session = requests.Session()
    seen: set[Tuple[str, str, str]] = set()
    rows: List[Dict] = []

    for year, month in month_range(start, end):
        for item in fetch_month(session, pol_ctr, pol, pol_name, pod_ctr, pod, pod_name, year, month):
            etd = item.get("etd", "")
            key = (item.get("vslNm", ""), item.get("voyNo", ""), etd)
            if key in seen:
                continue
            seen.add(key)
            if not (start.strftime("%Y%m%d") <= etd <= end.strftime("%Y%m%d")):
                continue
            rows.append(item)

    rows.sort(key=lambda r: (r.get("etd", ""), r.get("etdTm", "")))
    return rows


def fmt_dt(ymd: str, hm: str) -> str:
    """'20260901' + '2018' -> '2026-09-01 20:18'"""
    if not ymd or len(ymd) < 8:
        return ""
    date_part = f"{ymd[0:4]}-{ymd[4:6]}-{ymd[6:8]}"
    if hm and len(hm) == 4:
        return f"{date_part} {hm[0:2]}:{hm[2:4]}"
    return date_part


def build_excel(rows: List[Dict], pol_label: str, pod_label: str, out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "KMTC Schedule"

    headers = [
        "Vessel", "Voyage", "POL Terminal", "POL ETB", "POL ETD",
        "POD Terminal", "POD ETA", "Transit Time", "Booking Closing",
        "T/S", "Route Code", "Route Name",
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    for item in rows:
        close_time = item.get("closeTime", "") or ""
        ws.append([
            item.get("vslNm", ""),
            item.get("voyNo", ""),
            item.get("otrmlNm") or item.get("polTml", ""),
            fmt_dt(item.get("polEtb", ""), item.get("polEtbTm", "")),
            fmt_dt(item.get("etd", ""), item.get("etdTm", "")),
            item.get("itrmlNm") or item.get("podTml", ""),
            fmt_dt(item.get("eta", ""), item.get("etaTm", "")),
            (item.get("transitTime", "") or "").strip(),
            fmt_dt(close_time[:8], close_time[8:12]) if len(close_time) >= 12 else close_time,
            "Yes" if item.get("ts") == "Y" else "No",
            item.get("rteCd", ""),
            item.get("rteCdNm", ""),
        ])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    widths = [16, 9, 20, 16, 16, 24, 16, 16, 16, 6, 10, 34]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    # Title row above the header
    ws.insert_rows(1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=f"KMTC Schedule - {pol_label} to {pod_label}")
    title_cell.font = Font(bold=True, size=14, color="1F4E78")
    title_cell.alignment = Alignment(horizontal="left")

    wb.save(out_path)


def parse_year_month(s: str) -> dt.date:
    y, m = s.split("-")
    return dt.date(int(y), int(m), 1)


def main():
    parser = argparse.ArgumentParser(description="Export the KMTC vessel schedule between two ports to Excel.")
    parser.add_argument("--pol-ctr", default="TH", help="Loading port country code (default: TH)")
    parser.add_argument("--pol", default="LCH", help="Loading port code (default: LCH = Laem Chabang)")
    parser.add_argument("--pol-name", default="LAEM CHABANG", help="Loading port name as KMTC shows it")
    parser.add_argument("--pod-ctr", default="CN", help="Discharge port country code (default: CN)")
    parser.add_argument("--pod", default="SHA", help="Discharge port code (default: SHA = Shanghai)")
    parser.add_argument("--pod-name", default="SHANGHAI", help="Discharge port name as KMTC shows it")
    parser.add_argument("--start", default=None, help="Start month YYYY-MM (default: current month)")
    parser.add_argument("--end", default=None, help="End month YYYY-MM (default: December of current year)")
    parser.add_argument("--out", default=None, help="Output .xlsx path")
    args = parser.parse_args()

    today = dt.date.today()
    start = parse_year_month(args.start) if args.start else dt.date(today.year, today.month, 1)
    if args.end:
        end_month = parse_year_month(args.end)
        # last day of that month
        next_month = dt.date(end_month.year + (end_month.month == 12), (end_month.month % 12) + 1, 1)
        end = next_month - dt.timedelta(days=1)
    else:
        end = dt.date(today.year, 12, 31)

    print(f"Fetching KMTC schedule: {args.pol_name} ({args.pol}) -> {args.pod_name} ({args.pod})")
    print(f"Period: {start:%Y-%m} to {end:%Y-%m}")

    rows = collect_schedule(
        args.pol_ctr, args.pol, args.pol_name,
        args.pod_ctr, args.pod, args.pod_name,
        start, end,
    )
    print(f"Found {len(rows)} sailings.")

    out_path = args.out or f"KMTC Schedule ({args.pol}) ({args.pod}) {start:%Y-%m}_to_{end:%Y-%m}.xlsx"
    build_excel(rows, f"{args.pol_name} ({args.pol})", f"{args.pod_name} ({args.pod})", out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
