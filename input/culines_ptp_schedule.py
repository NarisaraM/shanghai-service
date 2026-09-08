#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
culines_ptp_schedule.py
========================

ดึงข้อมูล "Point to Point Schedule" จากเว็บไซต์ CU Lines
(https://www.culines.com/en/site/schedule_ptp)

ค่าเริ่มต้นของสคริปต์นี้ตั้งไว้ตามที่ขอ:
    Origin (ต้นทาง)       : Laem Chabang (แหลมฉบัง)
    Destination (ปลายทาง) : Shanghai
    Period (ช่วงเวลา)     : ตั้งแต่วันที่ 1 ของเดือนปัจจุบัน จนถึงวันที่ 31 ธันวาคม ของปีปัจจุบัน

วิธีทำงาน
---------
เว็บไซต์ culines.com เป็น Single Page Application ที่ดึงข้อมูลผ่าน endpoint
ภายใน 2 ตัว (เป็น public API ที่หน้าเว็บเรียกเอง ไม่ต้อง login/ไม่ต้องมี cookie):

1) ค้นหารหัสท่าเรือ (port code) จากชื่อท่าเรือ
   GET https://www.culines.com/search/getCommon
       ?address_url=https://eservice.culines.com/gnoss/CommonCodeGS.do
       &f_cmd=122
       &loc_nm=<ชื่อท่าเรือที่พิมพ์ในกล่อง Origin/Destination>
       &curl_type=get

2) ค้นหาตารางเดินเรือ (schedule) ระหว่าง 2 ท่าเรือ ในช่วงวันที่ที่กำหนด
   POST https://www.culines.com/search/getCommon
       address_url=https://eservice.culines.com/gnoss/CUP_HOM_3001GS.do
       curl_type=post
       f_cmd=3
       por_cd=<รหัสท่าเรือต้นทาง เช่น THLCH>
       del_cd=<รหัสท่าเรือปลายทาง เช่น CNSHA>
       frm_dt=<YYYY-MM-DD>
       to_dt=<YYYY-MM-DD>
       ts_ind=<"" = All, "D" = Direct, "T" = T/S (transshipment)>

สคริปต์นี้เรียก endpoint ทั้งสองตัวเหมือนที่หน้าเว็บทำ แล้วนำผลลัพธ์มาจัดเป็นตาราง
พร้อม export เป็นไฟล์ CSV และ Excel (.xlsx)

การติดตั้ง (ครั้งแรกครั้งเดียว)
--------------------------------
    pip install requests openpyxl

การใช้งาน
---------
    python culines_ptp_schedule.py
    python culines_ptp_schedule.py --origin "Laem Chabang" --destination "Shanghai"
    python culines_ptp_schedule.py --from 2026-09-01 --to 2026-12-31
    python culines_ptp_schedule.py --priority D          # D=Direct, T=T/S, (ไม่ใส่ = All)

หมายเหตุ
--------
* เว็บไซต์นี้เป็นบุคคลภายนอก (third-party) ปุ่ม/พารามิเตอร์อาจเปลี่ยนแปลงได้โดยไม่แจ้งล่วงหน้า
  หากสคริปต์ใช้งานไม่ได้ในอนาคต ให้เปิดหน้าเว็บ แล้วดู network request ของปุ่ม Search อีกครั้ง
* ตารางเดินเรือเป็นข้อมูลประมาณการ อาจเปลี่ยนแปลงได้ (ตามข้อความ disclaimer บนเว็บไซต์ต้นทาง)
"""

import argparse
import csv
import sys
from calendar import monthrange
from datetime import date

import requests

BASE_URL = "https://www.culines.com/search/getCommon"

# Address ปลายทางภายในของ culines ที่หน้าเว็บใช้เรียกตอนค้นหารหัสท่าเรือ / ตารางเดินเรือ
LOOKUP_ADDRESS = "https://eservice.culines.com/gnoss/CommonCodeGS.do"
SCHEDULE_ADDRESS = "https://eservice.culines.com/gnoss/CUP_HOM_3001GS.do"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Referer": "https://www.culines.com/en/site/schedule_ptp",
    "X-Requested-With": "XMLHttpRequest",
}


def get_port_code(session: requests.Session, location_name: str) -> dict:
    """แปลงชื่อท่าเรือ (เช่น 'Laem Chabang') เป็นรหัสท่าเรือ (เช่น 'THLCH')

    คืนค่า dict ของรายการแรกที่ตรงกัน เช่น
    {'locCd': 'THLCH', 'locNm': 'LAEM CHABANG', 'cntNm': 'THAILAND', ...}
    """
    params = {
        "address_url": LOOKUP_ADDRESS,
        "f_cmd": "122",
        "loc_nm": location_name,
        "curl_type": "get",
    }
    resp = session.get(BASE_URL, params=params, headers=HEADERS, timeout=30)
    resp.raise_for_status()
    data = resp.json()

    results = data.get("list") or []
    if not results:
        raise ValueError(
            f"ไม่พบท่าเรือที่ตรงกับ '{location_name}' บนเว็บไซต์ CU Lines "
            "ลองพิมพ์ชื่อท่าเรือให้ตรงกับที่กรอกในเว็บไซต์ (เช่นภาษาอังกฤษ)"
        )
    return results[0]


def fetch_schedule(
    session: requests.Session,
    por_cd: str,
    del_cd: str,
    frm_dt: str,
    to_dt: str,
    ts_ind: str = "",
) -> list:
    """เรียกตารางเดินเรือ point-to-point ระหว่างสองท่าเรือในช่วงวันที่ที่กำหนด

    ts_ind: "" = All, "D" = Direct, "T" = T/S (transshipment)
    คืนค่าเป็น list ของ dict แต่ละ dict คือ 1 เที่ยวเรือ (raw record จาก API)
    """
    payload = {
        "address_url": SCHEDULE_ADDRESS,
        "curl_type": "post",
        "f_cmd": "3",
        "por_cd": por_cd,
        "del_cd": del_cd,
        "frm_dt": frm_dt,
        "to_dt": to_dt,
        "ts_ind": ts_ind,
    }
    resp = session.post(BASE_URL, data=payload, headers=HEADERS, timeout=30)
    resp.raise_for_status()
    data = resp.json()

    if data.get("TRANS_RESULT_KEY") not in ("S", None):
        raise RuntimeError(f"เว็บไซต์ตอบกลับผิดพลาด: {data}")

    return data.get("list") or []


def normalize_record(rec: dict) -> dict:
    """แปลง record ดิบจาก API ให้เป็นคอลัมน์เดียวกับตารางที่แสดงบนหน้าเว็บ"""
    cct = (rec.get("cct") or "").split(".")[0]  # ตัดเศษวินาทีทิ้ง
    return {
        "Cargo Closing Time": f"{cct} {rec.get('cctDay', '')}".strip(),
        "Loading Port": rec.get("n1stLocNm", ""),
        "Departure Date": f"{rec.get('polEtdDt', '')} {rec.get('polEtdDay', '')}".strip(),
        "Discharging Port": rec.get("lstPodLocNm", ""),
        "Arrival Date": f"{rec.get('podEtaDt', '')} {rec.get('podEtaDay', '')}".strip(),
        "Lane": rec.get("n1stLaneNm") or rec.get("n1stLaneCd", ""),
        "Vessel": rec.get("n1stVslNm", ""),
        "Consortium/Voyage": rec.get("consVoyNr", ""),
        "T/Time Ocean (Day)": rec.get("ocnTzDys", ""),
        "T/Time Total (Day)": rec.get("ttlTzDys", ""),
        "Direct/T-S": "Direct" if rec.get("clInd") == "O" else rec.get("clInd", ""),
    }


def default_date_range() -> tuple:
    """ตั้งค่าเริ่มต้น: วันที่ 1 ของเดือนปัจจุบัน ถึง 31 ธันวาคม ปีปัจจุบัน"""
    today = date.today()
    frm = date(today.year, today.month, 1)
    to = date(today.year, 12, 31)
    return frm.isoformat(), to.isoformat()


def print_table(rows: list) -> None:
    if not rows:
        print("ไม่พบตารางเดินเรือในช่วงเวลาที่เลือก")
        return

    columns = list(rows[0].keys())
    widths = {c: max(len(c), *(len(str(r[c])) for r in rows)) for c in columns}

    header = " | ".join(c.ljust(widths[c]) for c in columns)
    print(header)
    print("-" * len(header))
    for r in rows:
        print(" | ".join(str(r[c]).ljust(widths[c]) for c in columns))


def save_csv(rows: list, path: str) -> None:
    if not rows:
        return
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def save_excel(rows: list, path: str, origin: str, destination: str, frm_dt: str, to_dt: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("(ข้ามการสร้างไฟล์ Excel: กรุณาติดตั้งด้วยคำสั่ง `pip install openpyxl`)")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "PTP Schedule"

    ws["A1"] = f"CU Lines Point to Point Schedule: {origin} -> {destination}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Period: {frm_dt} to {to_dt}"
    ws["A2"].font = Font(italic=True, size=10)

    header_row = 4
    if rows:
        columns = list(rows[0].keys())
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=col_name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F6F78")
            cell.alignment = Alignment(horizontal="center")

        for row_idx, rec in enumerate(rows, start=header_row + 1):
            for col_idx, col_name in enumerate(columns, start=1):
                ws.cell(row=row_idx, column=col_idx, value=rec[col_name])

        for col_idx, col_name in enumerate(columns, start=1):
            max_len = max(
                [len(col_name)] + [len(str(r[col_name])) for r in rows]
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4
    else:
        ws.cell(row=header_row, column=1, value="No sailings found for the selected period.")

    wb.save(path)


def main():
    parser = argparse.ArgumentParser(
        description="ดึงตารางเดินเรือ Point to Point Schedule จากเว็บไซต์ CU Lines"
    )
    parser.add_argument("--origin", default="Laem Chabang", help="ท่าเรือต้นทาง (default: Laem Chabang)")
    parser.add_argument("--destination", default="Shanghai", help="ท่าเรือปลายทาง (default: Shanghai)")
    parser.add_argument("--from", dest="frm_dt", default=None, help="วันที่เริ่มต้น YYYY-MM-DD (default: วันที่ 1 ของเดือนนี้)")
    parser.add_argument("--to", dest="to_dt", default=None, help="วันที่สิ้นสุด YYYY-MM-DD (default: 31 ธ.ค. ปีนี้)")
    parser.add_argument(
        "--priority",
        choices=["", "D", "T"],
        default="",
        help="All (default) / D = Direct / T = T/S (transshipment)",
    )
    parser.add_argument("--csv", dest="csv_path", default="culines_ptp_schedule.csv", help="path ไฟล์ CSV output")
    parser.add_argument("--xlsx", dest="xlsx_path", default="culines_ptp_schedule.xlsx", help="path ไฟล์ Excel output")
    args = parser.parse_args()

    default_frm, default_to = default_date_range()
    frm_dt = args.frm_dt or default_frm
    to_dt = args.to_dt or default_to

    session = requests.Session()

    print(f"กำลังค้นหารหัสท่าเรือสำหรับ '{args.origin}' และ '{args.destination}' ...")
    try:
        origin_info = get_port_code(session, args.origin)
        dest_info = get_port_code(session, args.destination)
    except Exception as exc:
        print(f"เกิดข้อผิดพลาดระหว่างค้นหารหัสท่าเรือ: {exc}", file=sys.stderr)
        sys.exit(1)

    por_cd = origin_info["locCd"]
    del_cd = dest_info["locCd"]
    print(
        f"Origin: {origin_info['locNm']} ({por_cd}) -> "
        f"Destination: {dest_info['locNm']} ({del_cd})"
    )
    print(f"Period: {frm_dt} to {to_dt}\n")

    try:
        raw_rows = fetch_schedule(session, por_cd, del_cd, frm_dt, to_dt, args.priority)
    except Exception as exc:
        print(f"เกิดข้อผิดพลาดระหว่างดึงตารางเดินเรือ: {exc}", file=sys.stderr)
        sys.exit(1)

    rows = [normalize_record(r) for r in raw_rows]

    # เรียงตามวันออกเรือ (Departure Date) จากเก่าไปใหม่
    rows.sort(key=lambda r: r["Departure Date"])

    print_table(rows)

    save_csv(rows, args.csv_path)
    save_excel(rows, args.xlsx_path, origin_info["locNm"], dest_info["locNm"], frm_dt, to_dt)

    print(f"\nพบทั้งหมด {len(rows)} เที่ยวเรือ")
    print(f"บันทึกไฟล์แล้ว: {args.csv_path}, {args.xlsx_path}")


if __name__ == "__main__":
    main()
